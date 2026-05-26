from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.conf import settings
from django.db.models import Sum, Count, F, ExpressionWrapper, DecimalField, Value
from django.db.models.functions import TruncMonth, Coalesce
from .models import Viaje, Unidad, TipoUnidad
from .forms import ViajeForm
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from .forms import AdminUserCreationForm
import requests as http_requests
from decimal import Decimal
import datetime

# --- Utilidades para reportes ---
from io import BytesIO
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Create your views here.

@login_required
def lista_viajes(request):
    """Vista para listar todos los viajes"""
    viajes = Viaje.objects.all()
    return render(request, 'trailers/lista_viajes.html', {'viajes': viajes})


@login_required
def agregar_viaje(request):
    """Vista para agregar un nuevo viaje"""
    if request.method == 'POST':
        form = ViajeForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Viaje agregado exitosamente.')
            return redirect('lista_viajes')
    else:
        form = ViajeForm()
    tipos_unidad = TipoUnidad.objects.all()
    return render(request, 'trailers/agregar_viaje.html', {'form': form, 'tipos_unidad': tipos_unidad})


@login_required
def editar_viaje(request, viaje_id):
    """Vista para editar un viaje existente"""
    viaje = get_object_or_404(Viaje, id=viaje_id)
    if request.method == 'POST':
        form = ViajeForm(request.POST, instance=viaje)
        if form.is_valid():
            form.save()
            messages.success(request, 'Viaje actualizado exitosamente.')
            return redirect('lista_viajes')
    else:
        form = ViajeForm(instance=viaje)
    tipos_unidad = TipoUnidad.objects.all()
    return render(request, 'trailers/editar_viaje.html', {'form': form, 'viaje': viaje, 'tipos_unidad': tipos_unidad})


@login_required
def detalle_viaje(request, viaje_id):
    """Vista de solo lectura para ver todos los datos de un viaje"""
    viaje = get_object_or_404(Viaje, id=viaje_id)
    return render(request, 'trailers/detalle_viaje.html', {'viaje': viaje})


@login_required
def borrar_viaje(request, viaje_id):
    """Vista para borrar un viaje"""
    viaje = get_object_or_404(Viaje, id=viaje_id)
    if request.method == 'POST':
        viaje.delete()
        messages.success(request, 'Viaje eliminado exitosamente.')
        return redirect('lista_viajes')
    return render(request, 'trailers/borrar_viaje.html', {'viaje': viaje})


def staff_check(user):
    return user.is_authenticated and user.is_staff


@login_required
def calcular_distancia(request):
    """Endpoint interno que consulta OpenRouteService y devuelve km entre origen y destino."""
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    origen = request.GET.get('origen', '').strip()
    destino = request.GET.get('destino', '').strip()

    if not origen or not destino:
        return JsonResponse({'error': 'Origen y destino son requeridos'}, status=400)

    api_key = settings.ORS_API_KEY
    if not api_key:
        return JsonResponse({'error': 'API key de ORS no configurada'}, status=500)

    def geocodificar(lugar):
        url = 'https://api.openrouteservice.org/geocode/search'
        params = {
            'api_key': api_key,
            'text': lugar + ', México',
            'size': 1,
            'boundary.country': 'MX',
        }
        resp = http_requests.get(url, params=params, timeout=10)
        resp.raise_for_status()
        features = resp.json().get('features', [])
        if not features:
            return None
        coords = features[0]['geometry']['coordinates']  # [lon, lat]
        return coords

    try:
        coords_origen = geocodificar(origen)
        coords_destino = geocodificar(destino)

        if not coords_origen or not coords_destino:
            return JsonResponse({'error': 'No se pudo encontrar alguna de las ubicaciones'}, status=404)

        url_ruta = 'https://api.openrouteservice.org/v2/directions/driving-car'
        headers = {'Authorization': api_key, 'Content-Type': 'application/json'}
        body = {'coordinates': [coords_origen, coords_destino]}
        resp_ruta = http_requests.post(url_ruta, json=body, headers=headers, timeout=15)
        resp_ruta.raise_for_status()
        data = resp_ruta.json()
        distancia_metros = data['routes'][0]['summary']['distance']
        km = round(distancia_metros / 1000, 2)
        return JsonResponse({'km': km})

    except http_requests.exceptions.RequestException as e:
        return JsonResponse({'error': f'Error al consultar ORS: {str(e)}'}, status=502)
    except (KeyError, IndexError):
        return JsonResponse({'error': 'Respuesta inesperada de ORS'}, status=502)


@login_required
@user_passes_test(staff_check)
def register(request):
    """Crear un nuevo usuario. Accesible solo para usuarios con is_staff."""
    if request.method == 'POST':
        form = AdminUserCreationForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Usuario creado correctamente.')
            return redirect('lista_viajes')
    else:
        form = AdminUserCreationForm()
    return render(request, 'trailers/register.html', {'form': form})


# =====================================================================
# REPORTES
# =====================================================================

MESES_ES = {
    1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
    5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
    9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre',
}

_CERO = Decimal('0')


def _get_anios_disponibles():
    """Años con al menos un viaje registrado."""
    anios = (
        Viaje.objects.filter(fecha_viaje__isnull=False)
        .dates('fecha_viaje', 'year')
    )
    return [d.year for d in anios] or [datetime.date.today().year]


def _calcular_datos_reporte(anio):
    """
    Devuelve un dict con:
      - resumen_mensual: lista de dicts por mes
      - resumen_unidades: lista de dicts por unidad
      - detalle_viajes: queryset de viajes del año con totales
      - totales_anuales: dict con sumas anuales
    """
    expr_opts = {'output_field': DecimalField(max_digits=14, decimal_places=2)}

    viajes_anio = (
        Viaje.objects.filter(fecha_viaje__year=anio, fecha_viaje__isnull=False)
        .annotate(
            total_gastos_calc=ExpressionWrapper(
                Coalesce(F('gastos_diesel'), _CERO)
                + Coalesce(F('gastos_casetas'), _CERO)
                + Coalesce(F('otros_gastos'), _CERO),
                **expr_opts
            ),
            utilidad_calc=ExpressionWrapper(
                Coalesce(F('costo_viaje'), _CERO)
                - Coalesce(F('gastos_diesel'), _CERO)
                - Coalesce(F('gastos_casetas'), _CERO)
                - Coalesce(F('otros_gastos'), _CERO),
                **expr_opts
            ),
        )
        .select_related('unidad', 'tipo_unidad')
        .order_by('fecha_viaje')
    )

    # --- Resumen mensual ---
    resumen_mensual_qs = (
        viajes_anio
        .annotate(mes=TruncMonth('fecha_viaje'))
        .values('mes')
        .annotate(
            num_viajes=Count('id'),
            ingresos=Coalesce(Sum('costo_viaje'), _CERO),
            diesel=Coalesce(Sum('gastos_diesel'), _CERO),
            casetas=Coalesce(Sum('gastos_casetas'), _CERO),
            otros=Coalesce(Sum('otros_gastos'), _CERO),
            total_gastos=Coalesce(
                Sum(ExpressionWrapper(
                    Coalesce(F('gastos_diesel'), _CERO)
                    + Coalesce(F('gastos_casetas'), _CERO)
                    + Coalesce(F('otros_gastos'), _CERO),
                    **expr_opts
                )), _CERO
            ),
            utilidad=Coalesce(
                Sum(ExpressionWrapper(
                    Coalesce(F('costo_viaje'), _CERO)
                    - Coalesce(F('gastos_diesel'), _CERO)
                    - Coalesce(F('gastos_casetas'), _CERO)
                    - Coalesce(F('otros_gastos'), _CERO),
                    **expr_opts
                )), _CERO
            ),
        )
        .order_by('mes')
    )

    resumen_mensual = []
    for fila in resumen_mensual_qs:
        fila['mes_nombre'] = MESES_ES.get(fila['mes'].month, '')
        fila['es_perdida'] = fila['utilidad'] < _CERO
        resumen_mensual.append(fila)

    # --- Resumen por unidad ---
    resumen_unidades_qs = (
        viajes_anio
        .values('unidad__numero_economico', 'unidad__marca', 'unidad__modelo')
        .annotate(
            num_viajes=Count('id'),
            ingresos=Coalesce(Sum('costo_viaje'), _CERO),
            diesel=Coalesce(Sum('gastos_diesel'), _CERO),
            casetas=Coalesce(Sum('gastos_casetas'), _CERO),
            otros=Coalesce(Sum('otros_gastos'), _CERO),
            total_gastos=Coalesce(
                Sum(ExpressionWrapper(
                    Coalesce(F('gastos_diesel'), _CERO)
                    + Coalesce(F('gastos_casetas'), _CERO)
                    + Coalesce(F('otros_gastos'), _CERO),
                    **expr_opts
                )), _CERO
            ),
            utilidad=Coalesce(
                Sum(ExpressionWrapper(
                    Coalesce(F('costo_viaje'), _CERO)
                    - Coalesce(F('gastos_diesel'), _CERO)
                    - Coalesce(F('gastos_casetas'), _CERO)
                    - Coalesce(F('otros_gastos'), _CERO),
                    **expr_opts
                )), _CERO
            ),
        )
        .order_by('-utilidad')
    )

    resumen_unidades = []
    for fila in resumen_unidades_qs:
        eco = fila.get('unidad__numero_economico') or '—'
        marca = fila.get('unidad__marca') or ''
        modelo = fila.get('unidad__modelo') or ''
        fila['unidad_label'] = f"{eco} {marca} {modelo}".strip()
        fila['es_perdida'] = fila['utilidad'] < _CERO
        resumen_unidades.append(fila)

    # --- Totales anuales ---
    totales = {
        'num_viajes': sum(r['num_viajes'] for r in resumen_mensual),
        'ingresos': sum(r['ingresos'] for r in resumen_mensual),
        'diesel': sum(r['diesel'] for r in resumen_mensual),
        'casetas': sum(r['casetas'] for r in resumen_mensual),
        'otros': sum(r['otros'] for r in resumen_mensual),
        'total_gastos': sum(r['total_gastos'] for r in resumen_mensual),
        'utilidad': sum(r['utilidad'] for r in resumen_mensual),
    }
    totales['es_perdida'] = totales['utilidad'] < _CERO

    return {
        'resumen_mensual': resumen_mensual,
        'resumen_unidades': resumen_unidades,
        'detalle_viajes': viajes_anio,
        'totales_anuales': totales,
    }


@login_required
def reportes(request):
    """Dashboard principal de reportes mensuales."""
    anios_disponibles = _get_anios_disponibles()
    anio_actual = datetime.date.today().year
    try:
        anio = int(request.GET.get('anio', anio_actual))
    except (ValueError, TypeError):
        anio = anio_actual

    datos = _calcular_datos_reporte(anio)

    return render(request, 'trailers/reportes.html', {
        'anio': anio,
        'anios_disponibles': anios_disponibles,
        **datos,
    })


@login_required
def exportar_reporte_excel(request):
    """Genera y descarga un archivo Excel con el reporte del año seleccionado."""
    try:
        anio = int(request.GET.get('anio', datetime.date.today().year))
    except (ValueError, TypeError):
        anio = datetime.date.today().year

    datos = _calcular_datos_reporte(anio)
    resumen_mensual = datos['resumen_mensual']
    resumen_unidades = datos['resumen_unidades']
    detalle_viajes = datos['detalle_viajes']
    totales = datos['totales_anuales']

    wb = openpyxl.Workbook()

    # --- Estilos comunes ---
    header_fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    subtotal_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    total_fill = PatternFill(start_color='0D1F3C', end_color='0D1F3C', fill_type='solid')
    perdida_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    ganancia_fill = PatternFill(start_color='375623', end_color='375623', fill_type='solid')
    white_font = Font(color='FFFFFF', bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    def _apply_header(ws, row_num, cols):
        for col, title in enumerate(cols, 1):
            cell = ws.cell(row=row_num, column=col, value=title)
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = thin_border

    def _apply_data_row(ws, row_num, values, fill=None, bold=False):
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
            if fill:
                cell.fill = fill
                cell.font = Font(color='FFFFFF', bold=bold)
            elif bold:
                cell.font = Font(bold=True)
            if isinstance(val, Decimal):
                cell.number_format = '"$"#,##0.00'
                cell.alignment = Alignment(horizontal='right')

    # ---- HOJA 1: Resumen Mensual ----
    ws1 = wb.active
    ws1.title = f'Resumen Mensual {anio}'
    ws1.merge_cells('A1:H1')
    titulo = ws1['A1']
    titulo.value = f'REPORTE MENSUAL {anio}'
    titulo.font = Font(size=14, bold=True, color='FFFFFF')
    titulo.fill = total_fill
    titulo.alignment = Alignment(horizontal='center')
    ws1.row_dimensions[1].height = 25

    encabezados = ['Mes', '# Viajes', 'Ingresos', 'Diesel', 'Casetas', 'Otros Gastos', 'Total Gastos', 'Utilidad / Pérdida']
    _apply_header(ws1, 2, encabezados)

    for i, fila in enumerate(resumen_mensual, 3):
        util_fill = perdida_fill if fila['es_perdida'] else ganancia_fill
        values = [
            fila['mes_nombre'], fila['num_viajes'],
            fila['ingresos'], fila['diesel'], fila['casetas'],
            fila['otros'], fila['total_gastos'], fila['utilidad'],
        ]
        _apply_data_row(ws1, i, values)
        util_cell = ws1.cell(row=i, column=8)
        util_cell.fill = util_fill
        util_cell.font = Font(color='FFFFFF', bold=True)

    # Fila de totales anuales
    total_row = len(resumen_mensual) + 3
    _apply_data_row(ws1, total_row, [
        f'TOTAL {anio}', totales['num_viajes'],
        totales['ingresos'], totales['diesel'], totales['casetas'],
        totales['otros'], totales['total_gastos'], totales['utilidad'],
    ], fill=subtotal_fill, bold=True)
    ws1.cell(row=total_row, column=8).fill = perdida_fill if totales['es_perdida'] else ganancia_fill

    col_widths = [14, 10, 14, 14, 12, 13, 14, 18]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ---- HOJA 2: Por Unidad ----
    ws2 = wb.create_sheet(f'Por Unidad {anio}')
    ws2.merge_cells('A1:H1')
    t2 = ws2['A1']
    t2.value = f'REPORTE POR UNIDAD {anio}'
    t2.font = Font(size=14, bold=True, color='FFFFFF')
    t2.fill = total_fill
    t2.alignment = Alignment(horizontal='center')
    ws2.row_dimensions[1].height = 25

    encabezados2 = ['Unidad', '# Viajes', 'Ingresos', 'Diesel', 'Casetas', 'Otros Gastos', 'Total Gastos', 'Utilidad / Pérdida']
    _apply_header(ws2, 2, encabezados2)

    for i, fila in enumerate(resumen_unidades, 3):
        util_fill = perdida_fill if fila['es_perdida'] else ganancia_fill
        values = [
            fila['unidad_label'], fila['num_viajes'],
            fila['ingresos'], fila['diesel'], fila['casetas'],
            fila['otros'], fila['total_gastos'], fila['utilidad'],
        ]
        _apply_data_row(ws2, i, values)
        ws2.cell(row=i, column=8).fill = util_fill
        ws2.cell(row=i, column=8).font = Font(color='FFFFFF', bold=True)

    col_widths2 = [28, 10, 14, 14, 12, 13, 14, 18]
    for i, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- HOJA 3: Detalle por Viaje ----
    ws3 = wb.create_sheet(f'Detalle Viajes {anio}')
    ws3.merge_cells('A1:L1')
    t3 = ws3['A1']
    t3.value = f'DETALLE DE VIAJES {anio}'
    t3.font = Font(size=14, bold=True, color='FFFFFF')
    t3.fill = total_fill
    t3.alignment = Alignment(horizontal='center')
    ws3.row_dimensions[1].height = 25

    encabezados3 = [
        '# Viaje', 'Fecha', 'Unidad', 'Origen', 'Destino',
        'Km', 'Ingresos', 'Diesel', 'Casetas', 'Otros', 'Total Gastos', 'Utilidad / Pérdida'
    ]
    _apply_header(ws3, 2, encabezados3)

    for i, v in enumerate(detalle_viajes, 3):
        unidad_label = str(v.unidad) if v.unidad else '—'
        util_fill = perdida_fill if v.utilidad_calc < _CERO else ganancia_fill
        values = [
            v.numero_viaje,
            v.fecha_viaje.strftime('%d/%m/%Y') if v.fecha_viaje else '—',
            unidad_label,
            v.origen, v.destino,
            float(v.km_distancia) if v.km_distancia else '—',
            v.costo_viaje,
            v.gastos_diesel, v.gastos_casetas, v.otros_gastos,
            v.total_gastos_calc, v.utilidad_calc,
        ]
        _apply_data_row(ws3, i, values)
        ws3.cell(row=i, column=12).fill = util_fill
        ws3.cell(row=i, column=12).font = Font(color='FFFFFF', bold=True)

    col_widths3 = [16, 12, 24, 20, 20, 10, 13, 13, 12, 11, 13, 18]
    for i, w in enumerate(col_widths3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # --- Respuesta HTTP ---
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_{anio}.xlsx"'
    return response


@login_required
def exportar_reporte_pdf(request):
    """Genera y descarga un PDF con el reporte del año seleccionado."""
    try:
        anio = int(request.GET.get('anio', datetime.date.today().year))
    except (ValueError, TypeError):
        anio = datetime.date.today().year

    datos = _calcular_datos_reporte(anio)
    resumen_mensual = datos['resumen_mensual']
    resumen_unidades = datos['resumen_unidades']
    detalle_viajes = datos['detalle_viajes']
    totales = datos['totales_anuales']

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=30
    )
    styles = getSampleStyleSheet()
    elements = []

    # Colores
    azul_oscuro = colors.HexColor('#1F3864')
    azul_medio = colors.HexColor('#2E75B6')
    verde = colors.HexColor('#375623')
    rojo = colors.HexColor('#C00000')
    gris_claro = colors.HexColor('#F2F2F2')

    title_style = ParagraphStyle('title', fontSize=16, textColor=colors.white,
                                  backColor=azul_oscuro, alignment=TA_CENTER,
                                  spaceAfter=6, spaceBefore=6, leading=20)
    section_style = ParagraphStyle('section', fontSize=12, textColor=colors.white,
                                    backColor=azul_medio, alignment=TA_CENTER,
                                    spaceAfter=4, spaceBefore=10, leading=16)
    normal = styles['Normal']

    def _fmt(val):
        if val is None:
            return '—'
        try:
            return f'${val:,.2f}'
        except Exception:
            return str(val)

    def _color_util(val):
        return verde if val >= _CERO else rojo

    elements.append(Paragraph(f'REPORTE ANUAL {anio}', title_style))
    elements.append(Spacer(1, 10))

    # --- Sección 1: Resumen Mensual ---
    elements.append(Paragraph('RESUMEN MENSUAL', section_style))

    col_widths_m = [70, 45, 75, 75, 65, 75, 80, 90]
    header_m = ['Mes', '# Viajes', 'Ingresos', 'Diesel', 'Casetas', 'Otros', 'Total Gastos', 'Utilidad / Pérdida']
    table_data = [header_m]
    table_style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), azul_oscuro),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, gris_claro]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]

    for idx, fila in enumerate(resumen_mensual, 1):
        row = [
            fila['mes_nombre'], str(fila['num_viajes']),
            _fmt(fila['ingresos']), _fmt(fila['diesel']),
            _fmt(fila['casetas']), _fmt(fila['otros']),
            _fmt(fila['total_gastos']), _fmt(fila['utilidad']),
        ]
        table_data.append(row)
        util_color = _color_util(fila['utilidad'])
        table_style_cmds += [
            ('BACKGROUND', (7, idx), (7, idx), util_color),
            ('TEXTCOLOR', (7, idx), (7, idx), colors.white),
            ('FONTNAME', (7, idx), (7, idx), 'Helvetica-Bold'),
        ]

    # Fila total
    total_row_idx = len(table_data)
    table_data.append([
        f'TOTAL {anio}', str(totales['num_viajes']),
        _fmt(totales['ingresos']), _fmt(totales['diesel']),
        _fmt(totales['casetas']), _fmt(totales['otros']),
        _fmt(totales['total_gastos']), _fmt(totales['utilidad']),
    ])
    table_style_cmds += [
        ('BACKGROUND', (0, total_row_idx), (6, total_row_idx), azul_medio),
        ('TEXTCOLOR', (0, total_row_idx), (6, total_row_idx), colors.white),
        ('FONTNAME', (0, total_row_idx), (-1, total_row_idx), 'Helvetica-Bold'),
        ('BACKGROUND', (7, total_row_idx), (7, total_row_idx), _color_util(totales['utilidad'])),
        ('TEXTCOLOR', (7, total_row_idx), (7, total_row_idx), colors.white),
    ]

    t_mensual = Table(table_data, colWidths=col_widths_m)
    t_mensual.setStyle(TableStyle(table_style_cmds))
    elements.append(t_mensual)
    elements.append(Spacer(1, 14))

    # --- Sección 2: Por Unidad ---
    elements.append(Paragraph('RESUMEN POR UNIDAD', section_style))

    col_widths_u = [130, 45, 75, 75, 65, 75, 80, 90]
    header_u = ['Unidad', '# Viajes', 'Ingresos', 'Diesel', 'Casetas', 'Otros', 'Total Gastos', 'Utilidad / Pérdida']
    table_data_u = [header_u]
    style_cmds_u = [
        ('BACKGROUND', (0, 0), (-1, 0), azul_oscuro),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, gris_claro]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]

    for idx, fila in enumerate(resumen_unidades, 1):
        row = [
            fila['unidad_label'], str(fila['num_viajes']),
            _fmt(fila['ingresos']), _fmt(fila['diesel']),
            _fmt(fila['casetas']), _fmt(fila['otros']),
            _fmt(fila['total_gastos']), _fmt(fila['utilidad']),
        ]
        table_data_u.append(row)
        util_color = _color_util(fila['utilidad'])
        style_cmds_u += [
            ('BACKGROUND', (7, idx), (7, idx), util_color),
            ('TEXTCOLOR', (7, idx), (7, idx), colors.white),
            ('FONTNAME', (7, idx), (7, idx), 'Helvetica-Bold'),
        ]

    t_unidades = Table(table_data_u, colWidths=col_widths_u)
    t_unidades.setStyle(TableStyle(style_cmds_u))
    elements.append(t_unidades)
    elements.append(Spacer(1, 14))

    # --- Sección 3: Detalle por viaje ---
    elements.append(Paragraph('DETALLE POR NÚMERO DE VIAJE', section_style))

    col_widths_d = [70, 52, 80, 70, 70, 38, 60, 60, 55, 50, 60, 80]
    header_d = ['# Viaje', 'Fecha', 'Unidad', 'Origen', 'Destino', 'Km', 'Ingresos', 'Diesel', 'Casetas', 'Otros', 'Total Gastos', 'Utilidad / Pérdida']
    table_data_d = [header_d]
    style_cmds_d = [
        ('BACKGROUND', (0, 0), (-1, 0), azul_oscuro),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ALIGN', (5, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (4, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, gris_claro]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('WORDWRAP', (0, 0), (-1, -1), True),
    ]

    for idx, v in enumerate(detalle_viajes, 1):
        unidad_label = str(v.unidad) if v.unidad else '—'
        km_str = f'{v.km_distancia:.1f}' if v.km_distancia else '—'
        row = [
            v.numero_viaje,
            v.fecha_viaje.strftime('%d/%m/%Y') if v.fecha_viaje else '—',
            unidad_label, v.origen, v.destino, km_str,
            _fmt(v.costo_viaje), _fmt(v.gastos_diesel),
            _fmt(v.gastos_casetas), _fmt(v.otros_gastos),
            _fmt(v.total_gastos_calc), _fmt(v.utilidad_calc),
        ]
        table_data_d.append(row)
        util_color = _color_util(v.utilidad_calc)
        style_cmds_d += [
            ('BACKGROUND', (11, idx), (11, idx), util_color),
            ('TEXTCOLOR', (11, idx), (11, idx), colors.white),
            ('FONTNAME', (11, idx), (11, idx), 'Helvetica-Bold'),
        ]

    t_detalle = Table(table_data_d, colWidths=col_widths_d)
    t_detalle.setStyle(TableStyle(style_cmds_d))
    elements.append(t_detalle)

    # Pie de página con fecha de generación
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width='100%', thickness=0.5, color=colors.grey))
    elements.append(Paragraph(
        f'Generado el {datetime.date.today().strftime("%d/%m/%Y")} — Sistema de Gestión de Trailers',
        ParagraphStyle('footer', fontSize=7, textColor=colors.grey, alignment=TA_CENTER)
    ))

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_{anio}.pdf"'
    return response

